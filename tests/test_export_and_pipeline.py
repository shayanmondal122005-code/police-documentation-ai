import io
import wave

import pytest
from docx import Document

from config import AI_DRAFT_WARNING, ENGINE_ASR, ENGINE_MOCK, ENGINE_PROMPTED
from documentation.clients import MockReportClient
from documentation.render import report_to_markdown
from documentation.report_generator import generate_report
from errors import (
    ConfigurationError,
    EmptyFileError,
    ExportError,
    FileTooLargeError,
    UnsupportedFileError,
)
from export.document_export import (
    OFFICER_REVIEW_HEADING,
    compose_document,
    export_docx,
    export_filename,
    export_markdown,
    export_pdf,
    export_txt,
)
from pipeline import PIPELINE_STEPS, ReportMetadata, build_engines, run_pipeline
from transcription.base import TranscriptionOptions
from utils.audio import detect_duration_seconds, format_duration, temporary_audio_file, validate_upload


def make_wav(seconds: float = 1.0, rate: int = 8000) -> bytes:
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(rate)
        wav.writeframes(b"\x00\x00" * int(rate * seconds))
    return buffer.getvalue()


@pytest.fixture
def report(mock_transcript, integrity):
    return generate_report(mock_transcript, integrity, MockReportClient(), duration="1:15", officer="SI Verma")


# ------------------------------------------------------------------ export


def test_compose_adds_warning_when_officer_deleted_it():
    document = compose_document("# Title\n\nBody only.")
    assert AI_DRAFT_WARNING in document
    assert f"## {OFFICER_REVIEW_HEADING}" in document
    assert "_No corrections entered._" in document


def test_compose_includes_officer_corrections_and_keeps_single_leading_warning(report):
    document = compose_document(report_to_markdown(report), "Time corrected to 20:00 after listening.")
    assert "Time corrected to 20:00 after listening." in document
    assert document.index(AI_DRAFT_WARNING) < document.index("Administrative Information")


def test_markdown_and_txt_exports_contain_required_content(report):
    document = compose_document(report_to_markdown(report), "ok")
    for data in (export_markdown(document), export_txt(document)):
        text = data.decode("utf-8")
        for required in ("POLICE FIELD INTERACTION REPORT", AI_DRAFT_WARNING, report.admin.recording_id,
                         report.admin.sha256, report.admin.processing_timestamp, "Officer Review / Corrections"):
            assert required.lower() in text.lower()
    assert "**" not in export_txt(document).decode("utf-8")


def test_docx_export_is_valid_and_contains_hindi_and_warning(report):
    document = compose_document(report_to_markdown(report), "ठीक है")
    data = export_docx(document)
    assert data[:2] == b"PK"
    parsed = Document(io.BytesIO(data))
    body = "\n".join(p.text for p in parsed.paragraphs)
    assert "POLICE FIELD INTERACTION REPORT" in body
    assert "सुनील यादव" in body
    assert report.admin.sha256 in body
    assert AI_DRAFT_WARNING in parsed.sections[0].header.paragraphs[0].text


def test_pdf_export_works_for_latin_text_and_includes_pdf_header():
    document = compose_document("# POLICE FIELD INTERACTION REPORT\n\n## Section\n\n- Item **bold** & <tag>\n")
    data = export_pdf(document)
    assert data.startswith(b"%PDF")


def test_pdf_export_refuses_devanagari_instead_of_producing_wrong_output(report):
    with pytest.raises(ExportError) as info:
        export_pdf(compose_document(report_to_markdown(report)))
    assert "DOCX" in info.value.user_message


def test_export_filename_is_sanitised():
    assert export_filename("REC-AB12/../x", "docx") == "police_report_REC-AB12x.docx"


def test_exports_never_contain_api_key(monkeypatch, report):
    monkeypatch.setenv("GEMINI_API_KEY", "AIza-SECRET-VALUE-123")
    document = compose_document(report_to_markdown(report))
    assert "AIza-SECRET-VALUE-123" not in document
    assert b"AIza-SECRET-VALUE-123" not in export_docx(document)


# ------------------------------------------------------------------ audio utilities


@pytest.mark.parametrize(
    "name, size, error",
    [("a.ogg", 10, UnsupportedFileError), ("noext", 10, UnsupportedFileError), ("a.wav", 0, EmptyFileError),
     ("a.wav", 10**10, FileTooLargeError)],
)
def test_validate_upload_rejects_bad_files(name, size, error):
    with pytest.raises(error):
        validate_upload(name, size)


@pytest.mark.parametrize("name", ["a.mp3", "a.WAV", "a.m4a", "a.mp4", "a.mpeg", "a.webm"])
def test_validate_upload_accepts_supported_formats(name):
    assert validate_upload(name, 100)


def test_wav_duration_detected_and_others_reported_unavailable():
    assert detect_duration_seconds(make_wav(2.0), "x.wav") == pytest.approx(2.0)
    assert detect_duration_seconds(b"not a wav", "x.wav") is None
    assert detect_duration_seconds(make_wav(), "x.mp3") is None
    assert format_duration(None) == "Unavailable" and format_duration(75) == "1:15"


def test_temp_audio_file_is_removed_after_use():
    with temporary_audio_file(b"data", ".wav") as path:
        assert path.read_bytes() == b"data"
    assert not path.exists()


def test_temp_audio_file_is_removed_even_on_error():
    with pytest.raises(RuntimeError):
        with temporary_audio_file(b"data", ".wav") as path:
            raise RuntimeError
    assert not path.exists()


# ------------------------------------------------------------------ pipeline


def test_mock_pipeline_runs_end_to_end_without_api_key(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    steps = []
    result = run_pipeline(
        "demo.wav", make_wav(), ENGINE_MOCK, TranscriptionOptions(), ReportMetadata(officer="SI Verma", case_reference="C-1"),
        progress=steps.append,
    )
    assert tuple(steps) == PIPELINE_STEPS
    assert result.integrity.recording_id.startswith("REC-")
    assert result.duration == "0:01"
    assert result.report.admin.officer == "SI Verma"
    assert result.transcript.raw_transcript in result.transcript.raw_transcript
    assert AI_DRAFT_WARNING in result.markdown


def test_pipeline_hash_is_of_the_uploaded_bytes():
    import hashlib

    data = make_wav()
    result = run_pipeline("a.wav", data, ENGINE_MOCK, TranscriptionOptions(), ReportMetadata())
    assert result.integrity.sha256 == hashlib.sha256(data).hexdigest()
    assert result.integrity.size_bytes == len(data)


def test_pipeline_rejects_invalid_uploads_before_any_processing():
    with pytest.raises(EmptyFileError):
        run_pipeline("a.wav", b"", ENGINE_MOCK, TranscriptionOptions(), ReportMetadata())
    with pytest.raises(UnsupportedFileError):
        run_pipeline("a.txt", b"x", ENGINE_MOCK, TranscriptionOptions(), ReportMetadata())


@pytest.mark.parametrize("engine", [ENGINE_PROMPTED, ENGINE_ASR])
def test_gemini_engines_require_api_key(monkeypatch, engine):
    monkeypatch.setattr("pipeline.get_api_key", lambda: None)
    with pytest.raises(ConfigurationError) as info:
        build_engines(engine)
    assert "GEMINI_API_KEY" in info.value.user_message


def test_unknown_engine_is_rejected(monkeypatch):
    monkeypatch.setattr("pipeline.get_api_key", lambda: "key")
    with pytest.raises(ConfigurationError):
        build_engines("whisper-not-installed")
