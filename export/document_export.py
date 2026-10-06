"""Export the (officer-edited) report as Markdown, TXT, DOCX or PDF.

All formats are produced from the same Markdown text, so edits made in the UI flow into
every export. Every export carries the AI-draft warning and an Officer Review section.
No secrets, logs or debug data are ever included.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from xml.sax.saxutils import escape

from config import AI_DRAFT_WARNING, HASH_DISCLAIMER, NOT_OFFICIAL_NOTICE, PRIVACY_NOTICE
from errors import ExportError

OFFICER_REVIEW_HEADING = "Officer Review / Corrections"
_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|(?<![A-Za-z0-9])_.+?_(?![A-Za-z0-9]))")
_DEVANAGARI = re.compile(r"[ऀ-ॿ]")


# ------------------------------------------------------------------ composition


def compose_document(report_markdown: str, officer_review: str = "") -> str:
    """Add the mandatory warning (if the officer removed it) and the Officer Review section."""
    body = report_markdown.strip()
    if AI_DRAFT_WARNING not in body:
        body = f"> **{AI_DRAFT_WARNING}**\n>\n> {NOT_OFFICIAL_NOTICE}\n\n{body}"
    review = officer_review.strip() or "_No corrections entered._"
    return (
        f"{body}\n\n"
        f"## {OFFICER_REVIEW_HEADING}\n\n"
        f"{review}\n\n"
        "- Reviewing officer: ______________________\n"
        "- Date reviewed: ______________________\n"
        "- Signature: ______________________\n\n"
        f"> **{AI_DRAFT_WARNING}**\n>\n> {HASH_DISCLAIMER}\n>\n> {PRIVACY_NOTICE}\n"
    )


def export_filename(recording_id: str, extension: str) -> str:
    """Safe download filename."""
    safe = re.sub(r"[^A-Za-z0-9_-]", "", recording_id) or "report"
    return f"police_report_{safe}.{extension}"


# ------------------------------------------------------------------ parsing


@dataclass
class Block:
    kind: str  # h1 | h2 | quote | bullet | numbered | para
    text: str
    indent: int = 0


def parse_blocks(markdown: str) -> list[Block]:
    """Parse the small Markdown subset produced by the renderer."""
    blocks: list[Block] = []
    for line in markdown.splitlines():
        stripped = line.rstrip()
        if not stripped.strip():
            continue
        content = stripped.strip()
        if stripped.startswith("# "):
            blocks.append(Block("h1", stripped[2:]))
        elif stripped.startswith("## "):
            blocks.append(Block("h2", stripped[3:]))
        elif stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            if text:
                blocks.append(Block("quote", text))
        elif re.match(r"^\s*- ", stripped):
            blocks.append(Block("bullet", content[2:], indent=1 if line.startswith(" ") else 0))
        elif re.match(r"^\d+\. ", stripped):
            blocks.append(Block("numbered", content))
        elif line.startswith(" ") and blocks:
            blocks[-1].text += "\n" + content
        else:
            blocks.append(Block("para", content))
    return blocks


def _tokens(text: str) -> list[tuple[str, str]]:
    """Split inline Markdown into (style, text) tokens: bold | italic | code | plain."""
    out: list[tuple[str, str]] = []
    pos = 0
    for match in _INLINE.finditer(text):
        if match.start() > pos:
            out.append(("plain", text[pos : match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            out.append(("bold", token[2:-2]))
        elif token.startswith("`"):
            out.append(("code", token[1:-1]))
        else:
            out.append(("italic", token[1:-1]))
        pos = match.end()
    if pos < len(text):
        out.append(("plain", text[pos:]))
    return out


def strip_inline(text: str) -> str:
    """Remove inline Markdown markers."""
    return "".join(token for _, token in _tokens(text))


# ------------------------------------------------------------------ formats


def export_markdown(document: str) -> bytes:
    return document.encode("utf-8")


def export_txt(document: str) -> bytes:
    """Plain-text rendering with headings underlined and quotes marked."""
    lines: list[str] = []
    for block in parse_blocks(document):
        text = strip_inline(block.text)
        if block.kind == "h1":
            lines += [text.upper(), "=" * len(text), ""]
        elif block.kind == "h2":
            lines += ["", text, "-" * len(text)]
        elif block.kind == "quote":
            lines.append(f">> {text}")
        elif block.kind == "bullet":
            lines.append(f"{'    ' * block.indent}- {text}")
        else:
            lines.append(text)
    return ("\n".join(lines) + "\n").encode("utf-8")


def export_docx(document: str) -> bytes:
    """Word document with the AI-draft warning in every page header."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor

        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:cs"), "Nirmala UI")  # complex-script (Devanagari) font

        header = doc.sections[0].header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        warn = header.add_run(AI_DRAFT_WARNING)
        warn.bold = True
        warn.font.color.rgb = RGBColor(0x9B, 0x1C, 0x1C)

        for block in parse_blocks(document):
            if block.kind == "h1":
                doc.add_heading(strip_inline(block.text), level=0)
                continue
            if block.kind == "h2":
                doc.add_heading(strip_inline(block.text), level=1)
                continue
            style = "List Bullet 2" if block.kind == "bullet" and block.indent else (
                "List Bullet" if block.kind == "bullet" else None
            )
            paragraph = doc.add_paragraph(style=style)
            if block.kind == "quote":
                paragraph.paragraph_format.left_indent = Pt(12)
            for line_no, line in enumerate(block.text.split("\n")):
                if line_no:
                    paragraph.add_run().add_break()
                for token_style, token in _tokens(line):
                    run = paragraph.add_run(token)
                    run.bold = token_style == "bold" or block.kind == "quote"
                    run.italic = token_style == "italic"
                    if token_style == "code":
                        run.font.name = "Consolas"
                    if block.kind == "quote":
                        run.font.color.rgb = RGBColor(0x9B, 0x1C, 0x1C)
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
    except Exception as exc:  # noqa: BLE001 - never leak internals to the UI
        raise ExportError("The DOCX export failed. Try the Markdown or TXT export instead.") from exc


def export_pdf(document: str) -> bytes:
    """PDF export (Latin text only).

    ReportLab cannot shape Devanagari conjuncts/vowel signs, so a PDF of Hindi/Bhojpuri text
    would be visibly wrong. Rather than produce a misleading record, such exports are refused
    and the officer is pointed to DOCX.
    """
    if _DEVANAGARI.search(document):
        raise ExportError(
            "PDF export is unavailable for reports that contain Hindi/Bhojpuri (Devanagari) text, "
            "because the PDF library cannot render it correctly. Use the DOCX export instead."
        )
    normalised = unicodedata.normalize("NFKC", document)
    try:
        normalised.encode("cp1252")
    except UnicodeEncodeError as exc:
        raise ExportError(
            "PDF export cannot render some characters in this report. Use the DOCX export instead."
        ) from exc
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        styles = getSampleStyleSheet()
        body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=13)
        quote = ParagraphStyle("quote", parent=body, textColor=colors.HexColor("#9B1C1C"), leftIndent=8)
        bullet = ParagraphStyle("bullet", parent=body, leftIndent=14, bulletIndent=4)
        bullet2 = ParagraphStyle("bullet2", parent=bullet, leftIndent=28, bulletIndent=18)

        def markup(text: str) -> str:
            parts = []
            for token_style, token in _tokens(text):
                safe = escape(token).replace("\n", "<br/>")
                parts.append(
                    {"bold": f"<b>{safe}</b>", "italic": f"<i>{safe}</i>", "code": f'<font name="Courier">{safe}</font>'}.get(
                        token_style, safe
                    )
                )
            return "".join(parts)

        story = []
        for block in parse_blocks(normalised):
            if block.kind == "h1":
                story.append(Paragraph(markup(block.text), styles["Title"]))
            elif block.kind == "h2":
                story += [Spacer(1, 6), Paragraph(markup(block.text), styles["Heading2"])]
            elif block.kind == "quote":
                story.append(Paragraph(f"<b>{markup(block.text)}</b>", quote))
            elif block.kind == "bullet":
                story.append(Paragraph(markup(block.text), bullet2 if block.indent else bullet, bulletText="•"))
            else:
                story.append(Paragraph(markup(block.text), body))

        def decorate(canvas, doc) -> None:  # noqa: ANN001 - ReportLab callback signature
            canvas.saveState()
            canvas.setFont("Helvetica-Bold", 8)
            canvas.setFillColor(colors.HexColor("#9B1C1C"))
            canvas.drawCentredString(A4[0] / 2, A4[1] - 10 * mm, AI_DRAFT_WARNING)
            canvas.setFillColor(colors.grey)
            canvas.drawCentredString(A4[0] / 2, 8 * mm, f"Page {doc.page}")
            canvas.restoreState()

        buffer = io.BytesIO()
        SimpleDocTemplate(
            buffer, pagesize=A4, topMargin=18 * mm, bottomMargin=16 * mm, leftMargin=18 * mm, rightMargin=18 * mm
        ).build(story, onFirstPage=decorate, onLaterPages=decorate)
        return buffer.getvalue()
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001 - never leak internals to the UI
        raise ExportError("The PDF export failed. Try the DOCX export instead.") from exc

